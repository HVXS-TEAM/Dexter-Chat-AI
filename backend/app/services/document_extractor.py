"""Text extraction for supported document formats."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Extract plain text from PDF, DOCX, PPTX, images, TXT, or MD."""
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1")
    elif suffix == ".pdf":
        from PyPDF2 import PdfReader

        reader = PdfReader(BytesIO(file_bytes))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document as DocxDocument

        document = DocxDocument(BytesIO(file_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        paragraphs.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        text = "\n\n".join(paragraphs)
    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(BytesIO(file_bytes))
        slides = []
        for slide in presentation.slides:
            slides.append("\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")))
        text = "\n\n".join(slides)
    elif suffix in {".png", ".jpg", ".jpeg"}:
        from PIL import Image
        import pytesseract

        text = pytesseract.image_to_string(Image.open(BytesIO(file_bytes)))
    else:
        raise ValueError(f"Unsupported document type: {suffix or 'unknown'}")

    cleaned = text.strip()
    if not cleaned:
        raise ValueError("The document does not contain extractable text.")
    return cleaned
